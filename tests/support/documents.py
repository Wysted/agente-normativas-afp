"""Tiny in-memory document builders for parser tests (no binary fixtures on disk)."""

import zipfile
from io import BytesIO

from docx import Document as new_docx
from pypdf import PdfReader, PdfWriter


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid PDF, one Helvetica text line per page (empty string = page without text)."""
    objects: list[bytes] = [b""] * (3 + 2 * len(pages))
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode()
    objects[2] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i, text in enumerate(pages):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode() if text else b""
        page_id = 4 + 2 * i
        objects[page_id - 1] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>"
        ).encode()
        objects[page_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def make_blank_pdf(page_count: int) -> bytes:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(blocks: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """`blocks` = (kind, text), kind in {"title", "h1", "h2", "p"}; optional table appended."""
    document = new_docx()
    for kind, text in blocks:
        if kind == "title":
            document.add_heading(text, level=0)
        elif kind in {"h1", "h2"}:
            document.add_heading(text, level=int(kind[1]))
        else:
            document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_encrypted_pdf(
    pages: list[str], user_password: str, owner_password: str = "owner", algorithm: str = "AES-256"
) -> bytes:
    """`make_pdf` encrypted; an empty user password means "opens without a password"."""
    writer = PdfWriter()
    writer.append_pages_from_reader(PdfReader(BytesIO(make_pdf(pages))))
    writer.encrypt(user_password=user_password, owner_password=owner_password, algorithm=algorithm)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx_with_merged_cells() -> bytes:
    """One row whose first two grid columns are merged (python-docx repeats that cell)."""
    document = new_docx()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(0, 2).text = "tail"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def with_extra_zip_member(docx: bytes, name: str, payload: bytes) -> bytes:
    """Append a deflated member to a DOCX: a tiny file that expands hugely (synthetic zip bomb)."""
    buffer = BytesIO(docx)
    with zipfile.ZipFile(buffer, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, payload)
    return buffer.getvalue()
